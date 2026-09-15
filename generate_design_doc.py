from docx import Document
from datetime import datetime
from pathlib import Path

output_path = Path(r"C:\Sysuser\Ishwar\AIProjects\Week3-Projects\MARKET_RESEARCH_AGENT\Market_Research_Agent_Design_Document.docx")
output_path.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
doc.add_heading('Market Research Agent Design Document', 0)
doc.add_paragraph(f'Generated on: {datetime.now().strftime("%Y-%m-%d")}')
doc.add_paragraph('This design document describes the architecture, workflow, and implementation of the Market Research Agent.')
doc.add_heading('1. Overview', 1)
doc.add_paragraph('The Market Research Agent provides a Streamlit-based interface for researching a target company, discovering competitors, gathering web and news evidence, and generating structured competitor summaries.')
doc.add_heading('2. Architecture', 1)
doc.add_paragraph('User input enters the Streamlit UI and is passed into a LangGraph workflow. The workflow performs discovery, evidence gathering, analyst synthesis, and queue management before the final reports are rendered.')
doc.add_paragraph('Main components: Streamlit UI, LangGraph workflow, You.com client, tool adapters, schema contracts, and security validation.')
doc.add_heading('3. Workflow', 1)
doc.add_paragraph('1. Validate user input.\n2. Discover candidate competitors.\n3. Research evidence for each competitor.\n4. Generate structured reports.\n5. Render results in the UI.')
doc.add_heading('4. Key Files', 1)
doc.add_paragraph('- app.py: Streamlit entry point')
doc.add_paragraph('- langgraph_workflow.py: orchestration logic')
doc.add_paragraph('- youcom_client.py: live API and local fallback behavior')
doc.add_paragraph('- youcom_tools.py: web and news tool adapters')
doc.add_paragraph('- schemas.py: report and evidence contracts')
doc.add_paragraph('- security.py: sanitization and validation')
doc.add_heading('5. Conclusion', 1)
doc.add_paragraph('The design emphasizes flexibility, deterministic offline operation, and a reviewable report pipeline, making it suitable for demos, evaluation, and extension.')

doc.save(output_path)
print(f"Saved to: {output_path}")
print(f"Exists: {output_path.exists()}")
